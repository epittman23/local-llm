import asyncio
import csv
import importlib
import logging
import os
import sys
import zipfile

import ftfy
import requests
from azure.identity import DefaultAzureCredential
from langchain_community.document_loaders import (
    AzureAIDocumentIntelligenceLoader,
    BSHTMLLoader,
    CSVLoader,
    Docx2txtLoader,
    PyPDFLoader,
    TextLoader,
)
from langchain_core.documents import Document

from local_llm.env import (
    AIOHTTP_CLIENT_SESSION_SSL,
    GLOBAL_LOG_LEVEL,
    MINERU_MAX_MARKDOWN_BYTES,
    REQUESTS_TIMEOUT_PROCESSING,
    REQUESTS_VERIFY,
)
from local_llm.retrieval.loaders.datalab_marker import DatalabMarkerLoader
from local_llm.retrieval.loaders.external_document import ExternalDocumentLoader
from local_llm.retrieval.loaders.mineru import MinerULoader
from local_llm.retrieval.loaders.mistral import MistralLoader
from local_llm.retrieval.loaders.paddleocr_vl import PADDLEOCR_VL_SUPPORTED_EXTENSIONS, PaddleOCRVLLoader
from local_llm.utils.headers import get_user_groups_for_custom_headers
from local_llm.utils.json_codec import JSONCodec

logging.basicConfig(stream=sys.stdout, level=GLOBAL_LOG_LEVEL)
log = logging.getLogger(__name__)

known_source_ext = [
    'go',
    'py',
    'java',
    'sh',
    'bat',
    'ps1',
    'cmd',
    'js',
    'ts',
    'css',
    'cpp',
    'hpp',
    'h',
    'c',
    'cs',
    'sql',
    'log',
    'ini',
    'pl',
    'pm',
    'r',
    'dart',
    'dockerfile',
    'env',
    'php',
    'hs',
    'hsc',
    'lua',
    'nginxconf',
    'conf',
    'm',
    'mm',
    'plsql',
    'perl',
    'rb',
    'rs',
    'db2',
    'scala',
    'bash',
    'swift',
    'vue',
    'svelte',
    'ex',
    'exs',
    'erl',
    'tsx',
    'jsx',
    'hs',
    'lhs',
    'json',
    'yaml',
    'yml',
    'toml',
]

known_archive_ext = {'docx', 'epub', 'odt', 'pptx', 'xlsx'}
known_archive_content_types = {
    'application/epub+zip',
    'application/vnd.oasis.opendocument.text',
    'application/vnd.openxmlformats-officedocument.presentationml.presentation',
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
}


DEFAULT_ARCHIVE_LIMIT = 100 * 1024 * 1024
DATALAB_MARKER_EXTENSIONS = [
    'pdf', 'xls', 'xlsx', 'ods', 'doc', 'docx', 'odt', 'ppt', 'pptx', 'odp',
    'html', 'epub', 'png', 'jpeg', 'jpg', 'webp', 'gif', 'tiff',
]  # fmt: skip
DOCX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
PPTX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
DOCUMENT_INTELLIGENCE_EXTENSIONS = ['pdf', 'docx', 'ppt', 'pptx']
DOCUMENT_INTELLIGENCE_CONTENT_TYPES = [DOCX_CONTENT_TYPE, 'application/vnd.ms-powerpoint', PPTX_CONTENT_TYPE]
EXCEL_CONTENT_TYPES = ['application/vnd.ms-excel', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet']
POWERPOINT_CONTENT_TYPES = ['application/vnd.ms-powerpoint', PPTX_CONTENT_TYPE]


def _community_loader(name: str):
    """A langchain_community loader class, looked up when first needed.

    Like a `from langchain_community.document_loaders import <name>` inside
    the function: ImportError when it can't be had, so callers can fall back.
    """
    try:
        return getattr(importlib.import_module('langchain_community.document_loaders'), name)
    except AttributeError as e:
        raise ImportError(f'cannot import name {name!r} from langchain_community.document_loaders') from e


def _unstructured_or_fallback(class_name: str, file_path: str, fallback, note: str, **kwargs):
    """An unstructured-based loader, or fallback() with a warning when the package isn't installed."""
    try:
        return _community_loader(class_name)(file_path, **kwargs)
    except ImportError:
        log.warning("The 'unstructured' package is not installed. " + note)
        return fallback()


def _unstructured_required(class_name: str, file_path: str, file_type: str, **kwargs):
    """An unstructured-based loader for a file type nothing else can read; ValueError without the package."""
    try:
        return _community_loader(class_name)(file_path, **kwargs)
    except ImportError:
        raise ValueError(
            f"Processing {file_type} files requires the 'unstructured' package. "
            'Install it with: pip install unstructured'
        )


class ExcelLoader:
    """Fallback Excel loader using pandas when unstructured is not installed."""

    def __init__(self, file_path):
        self.file_path = file_path

    def load(self) -> list[Document]:
        import pandas as pd

        text_parts = []
        xls = pd.ExcelFile(self.file_path)
        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name)
            text_parts.append(f'Sheet: {sheet_name}\n{df.to_string(index=False)}')
        return [
            Document(
                page_content='\n\n'.join(text_parts),
                metadata={'source': self.file_path},
            )
        ]


def get_csv_summary(filename: str, file_path: str, encoding: str) -> str | None:
    try:
        with open(file_path, newline='', encoding=encoding) as f:
            sample = f.read(4096)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample)
            except csv.Error:
                dialect = csv.excel

            total_rows = 0
            max_columns = 0
            headers = []
            for row in csv.reader(f, dialect):
                total_rows += 1
                max_columns = max(max_columns, len(row))
                if total_rows == 1:
                    headers = [header.lstrip('\ufeff') for header in row]
    except Exception:
        return None

    if total_rows == 0:
        return None

    return (
        f'Table: {total_rows} rows incl. header; '
        f'{max(total_rows - 1, 0)} data rows; '
        f'{max_columns} columns: {", ".join(headers)}.'
    )


class CSVLoaderWithSummary:
    def __init__(self, file_path: str, filename: str, encoding: str):
        self.file_path = file_path
        self.filename = filename
        self.encoding = encoding

    def load(self) -> list[Document]:
        docs = CSVLoader(self.file_path, encoding=self.encoding).load()
        if os.getenv('ENABLE_RAG_CSV_SUMMARY', 'False').lower() == 'true':
            summary = get_csv_summary(self.filename, self.file_path, self.encoding)
            if summary:
                docs.insert(0, Document(page_content=summary, metadata={'source': self.file_path, 'row': -1}))
        return docs


class PptxLoader:
    """Fallback PowerPoint loader using python-pptx when unstructured is not installed."""

    def __init__(self, file_path):
        self.file_path = file_path

    def load(self) -> list[Document]:
        from pptx import Presentation

        prs = Presentation(self.file_path)
        text_parts = []
        for i, slide in enumerate(prs.slides, 1):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    slide_texts.append(shape.text_frame.text)
            if slide_texts:
                text_parts.append(f'Slide {i}:\n' + '\n'.join(slide_texts))
        return [
            Document(
                page_content='\n\n'.join(text_parts),
                metadata={'source': self.file_path},
            )
        ]


class TikaLoader:
    def __init__(self, url, file_path, mime_type=None, extract_images=None, server_version='3'):
        self.url = url
        self.file_path = file_path
        self.mime_type = mime_type
        self.server_version = str(server_version or '3')

        self.extract_images = extract_images

    def load(self) -> list[Document]:
        with open(self.file_path, 'rb') as f:
            data = f.read()

        if self.mime_type is not None:
            headers = {'Content-Type': self.mime_type}
        else:
            headers = {}

        if self.extract_images is True:
            headers['X-Tika-PDFextractInlineImages'] = 'true'

        endpoint_path = 'tika/json/text' if self.server_version == '4' else 'tika/text'
        content_key = 'tk:content' if self.server_version == '4' else 'X-TIKA:content'
        endpoint = f'{self.url.rstrip("/")}/{endpoint_path}'

        r = requests.put(
            endpoint, data=data, headers=headers, verify=REQUESTS_VERIFY, timeout=REQUESTS_TIMEOUT_PROCESSING
        )

        if r.ok:
            raw_metadata = r.json()
            text = raw_metadata.get(content_key, '<No text content found>').strip()

            if 'Content-Type' in raw_metadata:
                headers['Content-Type'] = raw_metadata['Content-Type']

            log.debug('Tika extracted text: %s', text)

            return [Document(page_content=text, metadata=headers)]
        else:
            raise Exception(f'Error calling Tika: {r.reason}')


class DoclingLoader:
    def __init__(self, url, api_key=None, file_path=None, mime_type=None, params=None):
        self.url = url.rstrip('/')
        self.api_key = api_key
        self.file_path = file_path
        self.mime_type = mime_type

        self.params = params or {}

    def load(self) -> list[Document]:
        page_break_marker = '\f'
        with open(self.file_path, 'rb') as f:
            headers = {}
            if self.api_key:
                headers['X-Api-Key'] = f'{self.api_key}'

            r = requests.post(
                f'{self.url}/v1/convert/file',
                files={
                    'files': (
                        self.file_path,
                        f,
                        self.mime_type or 'application/octet-stream',
                    )
                },
                data={
                    'image_export_mode': 'placeholder',
                    'md_page_break_placeholder': page_break_marker,
                    # Keep Docling params as user-provided form values. Encoding nested
                    # values here would make Local LLM responsible for Docling's API
                    # quirks and could break when Docling changes its form contract.
                    **self.params,
                },
                headers=headers,
                verify=AIOHTTP_CLIENT_SESSION_SSL,
                timeout=REQUESTS_TIMEOUT_PROCESSING,
            )
        if r.ok:
            result = r.json()
            document_data = result.get('document', {})
            md_content = document_data.get('md_content', '')
            text = md_content or '<No text content found>'

            metadata = {'Content-Type': self.mime_type} if self.mime_type else {}
            if page_break_marker in md_content:
                documents = [
                    Document(page_content=page.strip(), metadata={**metadata, 'page': page_idx})
                    for page_idx, page in enumerate(md_content.split(page_break_marker))
                    if page.strip()
                ]
                if documents:
                    log.debug('Docling extracted text: %s', text)
                    return documents

            log.debug('Docling extracted text: %s', text)
            return [Document(page_content=text, metadata=metadata)]
        else:
            error_msg = f'Error calling Docling API: {r.reason}'
            if r.text:
                try:
                    error_data = r.json()
                    if 'detail' in error_data:
                        error_msg += f' - {error_data["detail"]}'
                except Exception:
                    error_msg += f' - {r.text}'
            raise Exception(f'Error calling Docling: {error_msg}')


class Loader:
    def __init__(self, engine: str = '', **kwargs):
        self.engine = engine
        self.user = kwargs.get('user', None)
        self.user_groups = kwargs.get('user_groups', None)
        self.metadata = kwargs.get('metadata', {})
        self.kwargs = kwargs

    def load(self, filename: str, file_content_type: str, file_path: str) -> list[Document]:
        loader = self._get_loader(filename, file_content_type, file_path)
        docs = loader.load()
        return [Document(page_content=ftfy.fix_text(doc.page_content), metadata=doc.metadata) for doc in docs]

    async def aload(self, filename: str, file_content_type: str, file_path: str) -> list[Document]:
        """
        Async wrapper around `load`.

        Document loaders dispatched by `_get_loader` (PyMuPDF, Unstructured,
        python-docx, Tika, etc.) are uniformly synchronous and CPU/IO-bound.
        Calling `load` directly from an async handler would block the event
        loop for the entire parse — minutes for large PDFs. This offloads
        the work to a worker thread so the loop stays responsive.
        """
        # Group lookup is async-only, so it must happen before `load`
        # is offloaded to a thread without a running event loop.
        if self.engine == 'external' and self.user_groups is None:
            self.user_groups = await get_user_groups_for_custom_headers(
                self.kwargs.get('EXTERNAL_DOCUMENT_LOADER_HEADERS'), self.user
            )

        return await asyncio.to_thread(self.load, filename, file_content_type, file_path)

    def _is_text_file(self, file_ext: str, file_content_type: str) -> bool:
        return file_ext in known_source_ext or (
            file_content_type
            and file_content_type.find('text/') >= 0
            # Avoid text/html files being detected as text
            and not file_content_type.find('html') >= 0
        )

    def _detect_text_encoding(self, file_path: str) -> str:
        """Detect the encoding of a text file with CJK-aware fallbacks.

        Langchain's ``TextLoader`` uses chardet internally when
        ``autodetect_encoding=True``, but chardet frequently misidentifies
        CJK encodings (e.g. GB18030 detected as GB2312 or even Cyrillic).
        This method replaces that by:

        1. Trying UTF-8 first (fast path for the vast majority of files).
        2. Using chardet as a *hint* to prioritise the right CJK codec
           family, but mapping subset names to their superset
           (e.g. GB2312 → gb18030).
        3. Validating that decoded text actually contains CJK characters,
           guarding against codecs that "succeed" but produce garbage.
        4. Falling back to latin-1 (always valid, ftfy fixes mojibake later).
        """
        try:
            with open(file_path, 'rb') as f:
                raw = f.read()
        except OSError:
            return 'utf-8'

        if not raw:
            return 'utf-8'

        # Fast path: most files are UTF-8
        try:
            raw.decode('utf-8')
            return 'utf-8'
        except UnicodeDecodeError as e:
            first_non_utf8 = e.start

        # Use chardet as a hint, not as ground truth
        import chardet

        # chardet is pure Python (~1.3s/MB), so sample around the first bad byte
        window = 256 * 1024
        sample_start = max(0, first_non_utf8 - window // 2)
        sample = raw[sample_start : sample_start + window]
        detected = chardet.detect(sample)
        # A stray byte can sit far from the real payload, leaving the sample with nothing to read
        if len(sample.translate(None, delete=bytes(range(128)))) < 64 and len(sample) < len(raw):
            detected = chardet.detect(raw)
        detected_enc = (detected.get('encoding') or '').lower().replace('-', '').replace('_', '')

        # Map chardet's detected encoding to the correct superset codec.
        # chardet often reports GB2312 for content that is actually GB18030;
        # GB18030 is a strict superset of both GB2312 and GBK.
        _ENC_FAMILY = {
            'gb2312': 'gb18030',
            'gb18030': 'gb18030',
            'gbk': 'gb18030',
            'big5': 'big5',
            'euckr': 'euc-kr',
            'eucjp': 'euc-jp',
            'iso2022jp': 'euc-jp',
            'shiftjis': 'shift_jis',
        }

        # Build priority list: chardet-hinted codec first, then remaining CJK
        base_order = ['gb18030', 'big5', 'euc-kr', 'euc-jp']
        hinted = _ENC_FAMILY.get(detected_enc)
        if hinted and hinted in base_order:
            ordered = [hinted] + [e for e in base_order if e != hinted]
        else:
            ordered = base_order

        for enc in ordered:
            try:
                text = raw.decode(enc)
                if text.strip() and self._has_cjk_characters(text):
                    log.info(
                        'Detected encoding %s for %s (chardet guessed %s)',
                        enc,
                        file_path,
                        detected.get('encoding'),
                    )
                    return enc
            except (UnicodeDecodeError, LookupError):
                continue

        # If chardet gave a non-CJK answer that isn't in our family map,
        # try it directly — it might be a valid Western encoding.
        chardet_encoding = detected.get('encoding')
        if chardet_encoding:
            try:
                raw.decode(chardet_encoding)
                log.info(
                    'Using chardet-detected encoding %s for %s',
                    chardet_encoding,
                    file_path,
                )
                return chardet_encoding
            except (UnicodeDecodeError, LookupError):
                pass

        # latin-1 is the ultimate fallback: every byte 0x00–0xFF is valid.
        # ftfy.fix_text() (applied downstream) repairs most mojibake that
        # results from treating Windows-1252 content as Latin-1.
        log.info('Falling back to latin-1 encoding for %s', file_path)
        return 'latin-1'

    @staticmethod
    def _has_cjk_characters(text: str, threshold: float = 0.05) -> bool:
        """Check if decoded text contains a meaningful proportion of CJK characters.

        This guards against codecs that technically "succeed" but decode the
        bytes into wrong Unicode codepoints (e.g. PUA chars, random symbols).
        A genuine CJK document should have at least ``threshold`` fraction of
        its non-whitespace characters in CJK Unicode blocks.
        """
        if not text:
            return False

        cjk_count = 0
        total = 0
        for ch in text:
            if ch.isspace():
                continue
            total += 1
            cp = ord(ch)
            if (
                0x4E00 <= cp <= 0x9FFF  # CJK Unified Ideographs
                or 0x3400 <= cp <= 0x4DBF  # CJK Extension A
                or 0x20000 <= cp <= 0x2A6DF  # CJK Extension B
                or 0x2A700 <= cp <= 0x2B73F  # CJK Extension C
                or 0x2B740 <= cp <= 0x2B81F  # CJK Extension D
                or 0xF900 <= cp <= 0xFAFF  # CJK Compatibility Ideographs
                or 0x3000 <= cp <= 0x303F  # CJK Symbols and Punctuation
                or 0x3040 <= cp <= 0x309F  # Hiragana
                or 0x30A0 <= cp <= 0x30FF  # Katakana
                or 0xAC00 <= cp <= 0xD7AF  # Hangul Syllables
                or 0xFF00 <= cp <= 0xFFEF  # Halfwidth and Fullwidth Forms
            ):
                cjk_count += 1

        if total == 0:
            return False

        return (cjk_count / total) >= threshold

    def _get_loader(self, filename: str, file_content_type: str, file_path: str):
        """The loader for this file: the configured engine's, when it applies, else one chosen by file type."""
        file_ext = filename.split('.')[-1].lower()
        if file_ext in known_archive_ext or file_content_type in known_archive_content_types:
            self._check_archive_size(file_path)

        engine_loader = {
            'external': self._external_loader,
            'tika': self._tika_loader,
            'datalab_marker': self._datalab_marker_loader,
            'docling': self._docling_loader,
            'document_intelligence': self._document_intelligence_loader,
            'mineru': self._mineru_loader,
            'mistral_ocr': self._mistral_ocr_loader,
            'paddleocr_vl': self._paddleocr_vl_loader,
        }.get(self.engine)
        loader = engine_loader(filename, file_ext, file_content_type, file_path) if engine_loader else None
        if loader is not None:
            return loader
        return self._file_type_loader(filename, file_ext, file_content_type, file_path)

    def _check_archive_size(self, file_path: str) -> None:
        """Rejects a zip-based document that decompresses to far more than its size (a zip bomb)."""
        max_file_size = self.kwargs.get('FILE_MAX_SIZE')
        try:
            max_file_size_bytes = int(max_file_size) * 1024 * 1024 if max_file_size else DEFAULT_ARCHIVE_LIMIT
        except (TypeError, ValueError):
            max_file_size_bytes = DEFAULT_ARCHIVE_LIMIT
        if max_file_size_bytes <= 0:
            return
        try:
            with zipfile.ZipFile(file_path) as archive:
                uncompressed_size = sum(entry.file_size for entry in archive.infolist())
        except (zipfile.BadZipFile, OSError):
            return
        max_bytes = min(max(10 * 1024 * 1024, os.path.getsize(file_path) * 100), max_file_size_bytes)
        if uncompressed_size > max_bytes:
            raise ValueError('Document archive is too large after decompression')

    def _text_loader(self, file_path: str):
        return TextLoader(file_path, encoding=self._detect_text_encoding(file_path))

    # Each engine's loader returns None when its configuration or the file
    # type doesn't apply; the file then goes to _file_type_loader.

    def _external_loader(self, filename, file_ext, file_content_type, file_path):
        if not (
            self.kwargs.get('EXTERNAL_DOCUMENT_LOADER_URL') and self.kwargs.get('EXTERNAL_DOCUMENT_LOADER_API_KEY')
        ):
            return None
        return ExternalDocumentLoader(
            file_path=file_path,
            url=self.kwargs.get('EXTERNAL_DOCUMENT_LOADER_URL'),
            api_key=self.kwargs.get('EXTERNAL_DOCUMENT_LOADER_API_KEY'),
            mime_type=file_content_type,
            user=self.user,
            user_groups=self.user_groups,
            headers=self.kwargs.get('EXTERNAL_DOCUMENT_LOADER_HEADERS'),
            metadata={**self.metadata, 'file_name': filename, 'file_content_type': file_content_type},
        )

    def _tika_loader(self, filename, file_ext, file_content_type, file_path):
        if not self.kwargs.get('TIKA_SERVER_URL'):
            return None
        if self._is_text_file(file_ext, file_content_type):
            return self._text_loader(file_path)
        return TikaLoader(
            url=self.kwargs.get('TIKA_SERVER_URL'),
            file_path=file_path,
            server_version=self.kwargs.get('TIKA_SERVER_VERSION'),
            extract_images=self.kwargs.get('PDF_EXTRACT_IMAGES'),
        )

    def _datalab_marker_loader(self, filename, file_ext, file_content_type, file_path):
        if not (self.kwargs.get('DATALAB_MARKER_API_KEY') and file_ext in DATALAB_MARKER_EXTENSIONS):
            return None
        api_base_url = self.kwargs.get('DATALAB_MARKER_API_BASE_URL', '')
        if not api_base_url or api_base_url.strip() == '':
            api_base_url = 'https://www.datalab.to/api/v1/marker'  # https://github.com/open-webui/open-webui/pull/16867#issuecomment-3218424349
        return DatalabMarkerLoader(
            file_path=file_path,
            api_key=self.kwargs['DATALAB_MARKER_API_KEY'],
            api_base_url=api_base_url,
            additional_config=self.kwargs.get('DATALAB_MARKER_ADDITIONAL_CONFIG'),
            use_llm=self.kwargs.get('DATALAB_MARKER_USE_LLM', False),
            skip_cache=self.kwargs.get('DATALAB_MARKER_SKIP_CACHE', False),
            force_ocr=self.kwargs.get('DATALAB_MARKER_FORCE_OCR', False),
            paginate=self.kwargs.get('DATALAB_MARKER_PAGINATE', False),
            strip_existing_ocr=self.kwargs.get('DATALAB_MARKER_STRIP_EXISTING_OCR', False),
            disable_image_extraction=self.kwargs.get('DATALAB_MARKER_DISABLE_IMAGE_EXTRACTION', False),
            format_lines=self.kwargs.get('DATALAB_MARKER_FORMAT_LINES', False),
            output_format=self.kwargs.get('DATALAB_MARKER_OUTPUT_FORMAT', 'markdown'),
        )

    def _docling_loader(self, filename, file_ext, file_content_type, file_path):
        if not self.kwargs.get('DOCLING_SERVER_URL'):
            return None
        if self._is_text_file(file_ext, file_content_type):
            return self._text_loader(file_path)
        params = self.kwargs.get('DOCLING_PARAMS', {})
        if not isinstance(params, dict):
            try:
                params = JSONCodec.loads(params)
            except JSONCodec.JSONDecodeError:
                log.error('Invalid DOCLING_PARAMS format, expected JSON object')
                params = {}
        return DoclingLoader(
            url=self.kwargs.get('DOCLING_SERVER_URL'),
            api_key=self.kwargs.get('DOCLING_API_KEY', None),
            file_path=file_path,
            mime_type=file_content_type,
            params=params,
        )

    def _document_intelligence_loader(self, filename, file_ext, file_content_type, file_path):
        # Only an empty endpoint disables it; an unset one (None) does not.
        if self.kwargs.get('DOCUMENT_INTELLIGENCE_ENDPOINT') == '' or not (
            file_ext in DOCUMENT_INTELLIGENCE_EXTENSIONS or file_content_type in DOCUMENT_INTELLIGENCE_CONTENT_TYPES
        ):
            return None
        if self.kwargs.get('DOCUMENT_INTELLIGENCE_KEY') != '':
            credentials = {'api_key': self.kwargs.get('DOCUMENT_INTELLIGENCE_KEY')}
        else:
            credentials = {'azure_credential': DefaultAzureCredential()}
        return AzureAIDocumentIntelligenceLoader(
            file_path=file_path,
            api_endpoint=self.kwargs.get('DOCUMENT_INTELLIGENCE_ENDPOINT'),
            **credentials,
            api_model=self.kwargs.get('DOCUMENT_INTELLIGENCE_MODEL'),
        )

    def _mineru_loader(self, filename, file_ext, file_content_type, file_path):
        if file_ext not in self.kwargs.get('MINERU_FILE_EXTENSIONS', ['pdf']):
            return None
        mineru_timeout = self.kwargs.get('MINERU_API_TIMEOUT', 300)
        if mineru_timeout:
            try:
                mineru_timeout = int(mineru_timeout)
            except ValueError:
                mineru_timeout = 300
        return MinerULoader(
            file_path=file_path,
            api_mode=self.kwargs.get('MINERU_API_MODE', 'local'),
            api_url=self.kwargs.get('MINERU_API_URL', 'http://localhost:8000'),
            api_key=self.kwargs.get('MINERU_API_KEY', ''),
            params=self.kwargs.get('MINERU_PARAMS', {}),
            timeout=mineru_timeout,
            max_markdown_bytes=MINERU_MAX_MARKDOWN_BYTES,
        )

    def _mistral_ocr_loader(self, filename, file_ext, file_content_type, file_path):
        # Mistral OCR currently only supports PDF and images. Only an empty
        # key disables it; an unset one (None) does not.
        if self.kwargs.get('MISTRAL_OCR_API_KEY') == '' or file_ext not in ['pdf']:
            return None
        return MistralLoader(
            base_url=self.kwargs.get('MISTRAL_OCR_API_BASE_URL'),
            api_key=self.kwargs.get('MISTRAL_OCR_API_KEY'),
            file_path=file_path,
            use_base64=self.kwargs.get('MISTRAL_OCR_USE_BASE64', False),
            user=self.user,
        )

    def _paddleocr_vl_loader(self, filename, file_ext, file_content_type, file_path):
        if not (
            self.kwargs.get('PADDLEOCR_VL_BASE_URL')
            and self.kwargs.get('PADDLEOCR_VL_TOKEN')
            and file_ext in PADDLEOCR_VL_SUPPORTED_EXTENSIONS
        ):
            return None
        return PaddleOCRVLLoader(
            api_url=self.kwargs.get('PADDLEOCR_VL_BASE_URL'),
            token=self.kwargs.get('PADDLEOCR_VL_TOKEN'),
            file_path=file_path,
        )

    def _file_type_loader(self, filename, file_ext, file_content_type, file_path):
        """The built-in loader for a file type, used when no engine applies; plain text for anything unknown."""
        if file_ext == 'pdf':
            return PyPDFLoader(
                file_path,
                extract_images=self.kwargs.get('PDF_EXTRACT_IMAGES'),
                mode=self.kwargs.get('PDF_LOADER_MODE', 'page'),
            )
        if file_ext == 'csv':
            return CSVLoaderWithSummary(file_path, filename, self._detect_text_encoding(file_path))
        if file_ext == 'rst':
            return _unstructured_or_fallback(
                'UnstructuredRSTLoader',
                file_path,
                lambda: self._text_loader(file_path),
                'Falling back to plain text loading for .rst file. Install it with: pip install unstructured',
                mode='elements',
            )
        if file_ext == 'xml':
            return _unstructured_or_fallback(
                'UnstructuredXMLLoader',
                file_path,
                lambda: self._text_loader(file_path),
                'Falling back to plain text loading for .xml file. Install it with: pip install unstructured',
            )
        if file_ext in ['htm', 'html']:
            return BSHTMLLoader(file_path, open_encoding='unicode_escape')
        if file_ext == 'md':
            return self._text_loader(file_path)
        office_loader = self._office_loader(file_ext, file_content_type, file_path)
        return office_loader if office_loader is not None else self._text_loader(file_path)

    def _office_loader(self, file_ext, file_content_type, file_path):
        """Loaders for e-books, Office and OpenDocument files, and email; None for anything else."""
        if file_content_type == 'application/epub+zip':
            return _unstructured_required('UnstructuredEPubLoader', file_path, '.epub')
        if file_content_type == DOCX_CONTENT_TYPE or file_ext == 'docx':
            return Docx2txtLoader(file_path)
        if file_ext == 'doc' or file_content_type == 'application/msword':
            return _unstructured_required('UnstructuredWordDocumentLoader', file_path, '.doc')
        if file_content_type in EXCEL_CONTENT_TYPES or file_ext in ['xls', 'xlsx']:
            return _unstructured_or_fallback(
                'UnstructuredExcelLoader',
                file_path,
                lambda: ExcelLoader(file_path),
                'Falling back to pandas for Excel file loading. '
                'Install unstructured for better results: pip install unstructured',
            )
        if file_content_type in POWERPOINT_CONTENT_TYPES or file_ext in ['ppt', 'pptx']:
            return _unstructured_or_fallback(
                'UnstructuredPowerPointLoader',
                file_path,
                lambda: PptxLoader(file_path),
                'Falling back to python-pptx for PowerPoint file loading. '
                'Install unstructured for better results: pip install unstructured',
            )
        if file_ext == 'msg':
            # unstructured parses .msg via python-oxmsg; avoids extract_msg's beautifulsoup4<4.14 conflict
            return _unstructured_required('UnstructuredEmailLoader', file_path, '.msg', process_attachments=False)
        if file_ext == 'odt':
            return _unstructured_required('UnstructuredODTLoader', file_path, '.odt')
        return None
